#!/usr/bin/env python3
# =============================================================================
# extrair-dados.py
# Recalcula os agregados publicáveis a partir das fontes primárias.
#
# USO
#   python3 scripts/extrair-dados.py --json <pasta-json-lattes> \
#                                    [--mapa <mapa_publicacoes_dados.xlsx>]
#
# EXEMPLO
#   python3 scripts/extrair-dados.py \
#     --json ~/Dropbox/codex/apcn/10_analises/entrega_atualizada_lattes_20260829/00_fontes_lattes_20260828/json \
#     --mapa ~/Dropbox/codex/apcn/10_analises/entrega_atualizada_lattes_20260829/01_publicacoes_docentes/mapa_publicacoes_dados.xlsx
#
# O QUE FAZ
#   Lê os currículos em JSON, deduplica artigos por DOI (e por título quando
#   não há DOI), cruza com o quartil Scopus quando o mapa é informado, e grava
#   data/extracao-<data>.json com os agregados.
#
# O QUE NÃO FAZ
#   NÃO escreve em assets/js/site-data.js. Isso é deliberado: todo número
#   publicado passa por conferência humana antes de entrar no portal. O script
#   imprime um bloco pronto para colar, e a decisão de publicar é da coordenação.
# =============================================================================

import argparse
import collections
import datetime
import glob
import json
import os
import re
import sys
import unicodedata

ANOS = list(range(2021, 2027))
EXCLUIR = {"GUILHERME JORGE BRIGOLINI SILVA"}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def limpa_doi(d):
    d = re.sub(r"^(https?://)?(dx\.)?doi\.org/", "", str(d or "").strip(), flags=re.I)
    return d.lower()


def carregar_curriculos(pasta):
    arquivos = sorted(glob.glob(os.path.join(pasta, "*.json")))
    if not arquivos:
        sys.exit(f"ERRO: nenhum .json em {pasta}")
    docs = []
    for p in arquivos:
        try:
            d = json.load(open(p, encoding="utf-8"))
        except Exception as e:
            print(f"  aviso: falha ao ler {os.path.basename(p)}: {e}", file=sys.stderr)
            continue
        nome = (d.get("informacoes_pessoais", {}).get("nome_completo") or "").strip()
        if nome.upper() in EXCLUIR:
            continue
        docs.append((os.path.basename(p), d))
    return docs


def quartis_do_mapa(caminho):
    """Devolve {codigo_periodico: quartil} e {chave_normalizada_revista: quartil}."""
    try:
        import openpyxl
    except ImportError:
        print("  aviso: openpyxl não instalado — quartis Scopus ignorados", file=sys.stderr)
        return {}, {}
    w = openpyxl.load_workbook(caminho, data_only=True)
    if "referencia_codigos" not in w.sheetnames:
        print("  aviso: aba 'referencia_codigos' ausente no mapa", file=sys.stderr)
        return {}, {}
    ws = w["referencia_codigos"]
    hdr = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    ix = {h: i for i, h in enumerate(hdr) if h}
    por_cod, por_rev = {}, {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        q = r[ix.get("quartil_scopus_2025", -1)] if "quartil_scopus_2025" in ix else None
        por_cod[r[0]] = q
        rev = r[ix["revista"]] if "revista" in ix else None
        if rev:
            por_rev[norm(rev)] = q
    return por_cod, por_rev


def main():
    ap = argparse.ArgumentParser(description="Recalcula agregados do portal REDEMAT.")
    ap.add_argument("--json", required=True, help="pasta com os currículos Lattes em JSON")
    ap.add_argument("--mapa", help="mapa_publicacoes_dados.xlsx (para quartis Scopus)")
    ap.add_argument("--out", default=None, help="arquivo de saída (padrão: data/extracao-<data>.json)")
    a = ap.parse_args()

    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    hoje = datetime.date.today().isoformat()
    saida = a.out or os.path.join(raiz, "data", f"extracao-{hoje}.json")

    print(f"Lendo currículos de {a.json}")
    docs = carregar_curriculos(a.json)
    print(f"  {len(docs)} currículos (excluídos: {', '.join(sorted(EXCLUIR))})")

    q_cod, q_rev = ({}, {})
    if a.mapa:
        print(f"Lendo quartis Scopus de {os.path.basename(a.mapa)}")
        q_cod, q_rev = quartis_do_mapa(a.mapa)
        print(f"  {sum(1 for v in q_cod.values() if v)} periódicos com quartil")

    unicos = {}          # chave -> (ano, quartil)
    por_docente = collections.Counter()
    com_doi = 0
    ori = collections.Counter()
    proj_total = 0
    paises = collections.Counter()

    for _, d in docs:
        arts = d.get("producao_bibliografica", {}).get("artigos_periodicos", []) or []
        for art in arts:
            try:
                ano = int(art.get("ano") or 0)
            except (TypeError, ValueError):
                continue
            if ano not in ANOS:
                continue
            por_docente[ano] += 1
            doi = limpa_doi(art.get("doi"))
            chave = doi or norm(art.get("titulo"))[:70]
            if not chave:
                continue
            if chave not in unicos:
                if doi:
                    com_doi += 1
                q = q_rev.get(norm(art.get("revista"))) or None
                unicos[chave] = (ano, q)

        o = d.get("orientacoes", {})
        for fase in ("concluidas", "em_andamento"):
            blo = o.get(fase, {}) or {}
            for nivel in ("mestrado", "doutorado", "pos_doutorado"):
                ori[f"{fase}_{nivel}"] += len(blo.get(nivel, []) or [])

        proj_total += len(d.get("projetos_pesquisa", []) or [])

        for ap_ in d.get("atuacao_profissional", []) or []:
            pais = (ap_.get("instituicao_pais") or "").strip()
            if pais and pais.lower() not in ("brasil", "brazil"):
                paises[pais] += 1

    por_ano = collections.Counter(v[0] for v in unicos.values())
    por_q = collections.Counter(v[1] or "sem_quartil" for v in unicos.values())

    res = {
        "gerado_em": datetime.datetime.now().isoformat(timespec="seconds"),
        "fonte_json": a.json,
        "fonte_mapa": a.mapa,
        "n_docentes": len(docs),
        "producao": {
            "total_unico": len(unicos),
            "total_somado": sum(por_docente.values()),
            "com_doi": com_doi,
            "unicos_por_ano": dict(sorted(por_ano.items())),
            "somado_por_ano": dict(sorted(por_docente.items())),
            "quartil_unicos": dict(por_q),
        },
        "orientacoes": dict(ori),
        "projetos_declarados": proj_total,
        "paises": dict(paises.most_common()),
    }

    os.makedirs(os.path.dirname(saida), exist_ok=True)
    json.dump(res, open(saida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print(f"\nEscrito: {saida}")
    print("\n" + "=" * 68)
    print("BLOCO PARA CONFERIR E COLAR EM assets/js/site-data.js")
    print("=" * 68)
    p = res["producao"]
    print(f"""
    unicos:      {{ {', '.join(f'{k}: {v}' for k, v in p['unicos_por_ano'].items())} }},
    por_docente: {{ {', '.join(f'{k}: {v}' for k, v in p['somado_por_ano'].items())} }},
    total_unico: {p['total_unico']},
    total_somado: {p['total_somado']},
    com_doi: {p['com_doi']},
""")
    q = p["quartil_unicos"]
    print(f"""    scopus.unicos: {{ Q1: {q.get('Q1', 0)}, Q2: {q.get('Q2', 0)}, Q3: {q.get('Q3', 0)}, Q4: {q.get('Q4', 0)}, sem_quartil: {q.get('sem_quartil', 0)} }},
""")
    print("=" * 68)
    print("""
IMPORTANTE — este script NÃO altera o portal.

Antes de publicar qualquer número alterado:
  1. Compare com a versão anterior e explique cada variação.
  2. Confira os totais na Plataforma Sucupira.
  3. Atualize o campo `coleta` da fonte correspondente em meta.fontes.
  4. Se um número mudou de status (PRELIMINAR -> VALIDADO), registre em
     data/data_conflicts.csv o que passou a sustentá-lo.
""")


if __name__ == "__main__":
    main()
