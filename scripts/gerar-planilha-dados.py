#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera planilhas/REDEMAT-dados-do-site.xlsx — a planilha de manutenção.

Para que serve
--------------
Quem mantém o Programa não edita JavaScript. Esta planilha é a porta de entrada
dos dados que mudam com frequência: docentes (com os sites pessoal e de grupo),
disciplinas e ementas. Depois de editada, `scripts/planilha-para-dados.py`
devolve o conteúdo para os arquivos que o portal lê.

Regra: a planilha é gerada A PARTIR dos dados atuais, nunca do zero. Rodar este
script de novo sobrevive a edições — ele relê o estado corrente e remonta.
Edições feitas na planilha e ainda não aplicadas, porém, são perdidas: aplique
antes de regerar.
"""
import json, os, re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, 'planilhas', 'REDEMAT-dados-do-site.xlsx')

FONTE = 'Arial'
CAB_FILL = PatternFill('solid', fgColor='1B2A4A')
CAB_FONT = Font(name=FONTE, size=10, bold=True, color='FFFFFF')
EDIT_FILL = PatternFill('solid', fgColor='FFF7CC')   # amarelo: pode editar
LEITURA_FONT = Font(name=FONTE, size=10, color='666666')
NORMAL = Font(name=FONTE, size=10)
BORDA = Border(*[Side('thin', color='D7DCE5')] * 4)


def le_historico():
    p = os.path.join(RAIZ, 'assets', 'js', 'historico-dados.js')
    s = open(p, encoding='utf-8').read()
    return json.loads(re.search(r'=\s*(\{.*\});?\s*$', s, re.S).group(1))


def le_json(nome):
    p = os.path.join(RAIZ, 'data', nome)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else None


def le_sites():
    """Sites pessoais e de grupo já registrados. Arquivo separado de propósito:
    o histórico vem da coleta CAPES e é regerado por script; o site pessoal é
    declaração da pessoa e não pode ser apagado por uma regeração."""
    d = le_json('sites-docentes.json')
    return {x['nome']: x for x in d['docentes']} if d else {}


def escreve(ws, cabecalhos, larguras, linhas, editaveis):
    ws.freeze_panes = 'A2'
    for c, (t, w) in enumerate(zip(cabecalhos, larguras), 1):
        cel = ws.cell(1, c, t)
        cel.fill, cel.font = CAB_FILL, CAB_FONT
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 30
    for r, linha in enumerate(linhas, 2):
        for c, v in enumerate(linha, 1):
            cel = ws.cell(r, c, v)
            cel.font = NORMAL if cabecalhos[c - 1] in editaveis else LEITURA_FONT
            cel.alignment = Alignment(vertical='top', wrap_text=True)
            cel.border = BORDA
            if cabecalhos[c - 1] in editaveis:
                cel.fill = EDIT_FILL
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cabecalhos)), len(linhas) + 1)


def main():
    h = le_historico()
    sites = le_sites()
    disc = le_json('disciplinas.json')

    wb = Workbook()

    # ---------------------------------------------------------------- Leia-me
    ws = wb.active
    ws.title = 'Leia-me'
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 96
    linhas_txt = [
        ('REDEMAT — planilha de manutenção do portal', ''),
        ('', ''),
        ('Como usar', 'Edite apenas as células de fundo amarelo. As de texto cinza vêm de '
                      'outra fonte (coleta CAPES, Lattes, proposta APCN) e são sobrescritas '
                      'na próxima atualização automática — alterá-las aqui não tem efeito.'),
        ('Depois de editar', 'Salve o arquivo e avise: o script scripts/planilha-para-dados.py '
                             'devolve o conteúdo para os arquivos que o portal lê. Enquanto '
                             'isso não for feito, o site continua mostrando o conteúdo antigo.'),
        ('Linha de exemplo', 'A primeira linha de dados de cada aba traz valores reais, não '
                             'fictícios. Não existe linha de exemplo a apagar.'),
        ('Não publicar', 'Nunca inclua CPF, telefone pessoal, endereço residencial, e-mail de '
                         'discente ou qualquer pontuação individual por docente. A planilha é '
                         'insumo de um site público.'),
        ('', ''),
        ('Abas', ''),
        ('Docentes', 'Quadro atual. As colunas de site pessoal, site do grupo e nome do grupo '
                     'são para preencher — o portal só mostra o que estiver aqui.'),
        ('Disciplinas', 'Estrutura curricular: nome, carga, créditos, ementa, referências e '
                        'docentes responsáveis.'),
        ('', ''),
        ('Conferência', ''),
        ('Docentes na planilha', '=COUNTA(Docentes!A2:A200)'),
        ('Com site pessoal informado', '=COUNTA(Docentes!F2:F200)'),
        ('Com site de grupo informado', '=COUNTA(Docentes!G2:G200)'),
        ('Disciplinas na planilha', '=COUNTA(Disciplinas!A2:A200)'),
        ('Soma de créditos do catálogo', '=SUM(Disciplinas!D2:D200)'),
    ]
    for r, (a, b) in enumerate(linhas_txt, 1):
        ca, cb = ws.cell(r, 1, a), ws.cell(r, 2, b)
        ca.font = Font(name=FONTE, size=11, bold=(r == 1 or a in ('Abas', 'Conferência')))
        cb.font = NORMAL
        cb.alignment = Alignment(vertical='top', wrap_text=True)
    ws['A1'].font = Font(name=FONTE, size=14, bold=True, color='1B2A4A')

    # --------------------------------------------------------------- Docentes
    cab = ['Nome', 'ID Lattes', 'IES', 'Categoria', 'No quadro',
           'Site pessoal', 'Site do grupo', 'Nome do grupo', 'Observação']
    editaveis = {'Site pessoal', 'Site do grupo', 'Nome do grupo', 'Observação'}
    quadro = [d for d in h['docentes'] if d.get('no_quadro')]
    quadro.sort(key=lambda d: d['nome'])
    linhas = []
    for d in quadro:
        s = sites.get(d['nome'], {})
        linhas.append([d['nome'], d.get('lattes') or '', ', '.join(d.get('ies') or []),
                       d.get('categoria') or '', 'sim' if d.get('no_quadro') else 'não',
                       s.get('site_pessoal', ''), s.get('site_grupo', ''),
                       s.get('nome_grupo', ''), s.get('observacao', '')])
    escreve(wb.create_sheet('Docentes'), cab, [34, 18, 10, 14, 10, 34, 30, 24, 30],
            linhas, editaveis)

    # ------------------------------------------------------------ Disciplinas
    cab = ['Código', 'Nome', 'Horas', 'Créditos', 'Grupo', 'Área', 'Linha',
           'Docentes responsáveis', 'Ementa', 'Referências', 'Status']
    editaveis = {'Nome', 'Horas', 'Créditos', 'Docentes responsáveis',
                 'Ementa', 'Referências'}
    linhas = []
    for x in (disc or {}).get('disciplinas', []):
        linhas.append([x['id'], x['nome'], x['horas'], x['creditos'], x['grupo'],
                       x.get('area') or '', x.get('linha') or '',
                       '; '.join(x.get('docentes') or []),
                       x.get('ementa') or '', x.get('referencias') or '',
                       x.get('status') or ''])
    escreve(wb.create_sheet('Disciplinas'), cab,
            [9, 42, 8, 9, 18, 24, 26, 34, 70, 60, 12], linhas, editaveis)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)
    print('gravado:', SAIDA)
    print('  docentes no quadro: %d' % len(quadro))
    print('  disciplinas: %d' % len(linhas))


if __name__ == '__main__':
    main()
